from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

import analyze_completed_audit as analysis
from build_audit_package import CODING_COLUMNS, ENUMS, ROOT


def inspect_workbook(path: Path, expected_rows: int) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    assert wb.sheetnames == ["00_READ_ME", "01_REVIEW"]
    ws = wb["01_REVIEW"]
    headers = [c.value for c in ws[1]]
    assert ws.max_row == expected_rows + 1
    assert len(set(ws.cell(r, 1).value for r in range(2, ws.max_row + 1))) == expected_rows
    forbidden = {"topic", "family_id_norm", "hmm", "stm", "risk_level"}
    assert not forbidden.intersection(headers)
    for col in CODING_COLUMNS:
        idx = headers.index(col) + 1
        assert all(ws.cell(r, idx).value in (None, "") for r in range(2, ws.max_row + 1))
    assert len(ws.data_validations.dataValidation) >= len(ENUMS)
    assert ws.protection.sheet


def fill_synthetic(src: Path, dst: Path) -> None:
    shutil.copy2(src, dst)
    wb = openpyxl.load_workbook(dst)
    ws = wb["01_REVIEW"]
    headers = [c.value for c in ws[1]]
    for row in range(2, ws.max_row + 1):
        for col, allowed in ENUMS.items():
            idx = headers.index(col) + 1
            if col == "review_status":
                value = "COMPLETE"
            else:
                value = allowed[(row - 2) % min(2, len(allowed))]
            ws.cell(row, idx, value)
        ws.cell(row, headers.index("exact_evidence_quote") + 1, "Synthetic quote for package test only")
        ws.cell(row, headers.index("decision_reason") + 1, "Synthetic reason for package test only")
    wb.save(dst)


def main() -> None:
    r1 = ROOT / "01_round1_blinded_review.xlsx"
    r2 = ROOT / "02_round2_delayed_retest.xlsx"
    key_path = ROOT / "99_DO_NOT_OPEN_topic_key.xlsx"
    inspect_workbook(r1, 87)
    inspect_workbook(r2, 18)

    key = pd.read_excel(key_path, sheet_name="DO_NOT_OPEN")
    assert key["review_id"].nunique() == 87
    assert key["family_id_norm"].nunique() == 87
    assert key["topic"].value_counts().to_dict() == {22: 49, 25: 38}
    assert key["retest_id"].notna().sum() == 18
    assert key.loc[key["retest_id"].notna(), "review_id"].nunique() == 18

    with tempfile.TemporaryDirectory(prefix="topic_audit_test_") as tmp:
        tmp_path = Path(tmp)
        test_r1 = tmp_path / "r1.xlsx"
        test_r2 = tmp_path / "r2.xlsx"
        fill_synthetic(r1, test_r1)
        fill_synthetic(r2, test_r2)
        analysis.ROUND1 = test_r1
        analysis.ROUND2 = test_r2
        analysis.KEY = key_path
        analysis.OUT_XLSX = tmp_path / "results.xlsx"
        analysis.OUT_MD = tmp_path / "results.md"
        analysis.main()
        assert analysis.OUT_XLSX.exists()
        assert analysis.OUT_MD.exists()
        result_wb = openpyxl.load_workbook(analysis.OUT_XLSX, data_only=True)
        assert result_wb.sheetnames == [
            "01_Descriptive", "02_Topic_Comparison", "03_Intra_Rater", "04_Unblinded_Data"
        ]
    print("PASS: workbook structure, blinding, counts, validations, and synthetic analysis")


if __name__ == "__main__":
    main()
