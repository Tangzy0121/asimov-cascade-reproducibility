from __future__ import annotations

import re
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from analyze_completed_audit import METRICS, ROOT, outcome_table


INPUT = ROOT / "audit_results.xlsx"
OUTPUT = ROOT / "duplicate_text_sensitivity.xlsx"


def normalize_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def collapse_exact_duplicates(data: pd.DataFrame):
    """Collapse within-topic rows connected by identical claim or abstract.

    This is an exploratory sensitivity analysis prompted by duplicate source
    text detected while the topic key was still unopened. The preregistered
    87-row analysis remains the primary analysis.
    """
    parent = list(range(len(data)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(a: int, b: int) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra

    match_fields: dict[tuple[int, int], set[str]] = defaultdict(set)
    for field in ("first_claim_english", "abstract_english"):
        groups: dict[tuple[int, str], list[int]] = defaultdict(list)
        for i, row in data.iterrows():
            text = normalize_text(row[field])
            if text:
                groups[(int(row["topic"]), text)].append(i)
        for indices in groups.values():
            if len(indices) < 2:
                continue
            for i in indices[1:]:
                union(indices[0], i)
            for a in indices:
                for b in indices:
                    if a < b:
                        match_fields[(a, b)].add(field)

    clusters: dict[int, list[int]] = defaultdict(list)
    for i in range(len(data)):
        clusters[find(i)].append(i)

    keep = []
    group_rows = []
    duplicate_group_number = 0
    for indices in clusters.values():
        chosen = min(indices, key=lambda i: str(data.loc[i, "record_id"]))
        keep.append(chosen)
        if len(indices) == 1:
            continue
        duplicate_group_number += 1
        fields = set()
        for a in indices:
            for b in indices:
                if a < b:
                    fields.update(match_fields.get((a, b), set()))
        group_rows.append(
            {
                "duplicate_group": f"D{duplicate_group_number:02d}",
                "topic": int(data.loc[chosen, "topic"]),
                "group_size": len(indices),
                "retained_record_id": data.loc[chosen, "record_id"],
                "record_ids": "; ".join(sorted(str(data.loc[i, "record_id"]) for i in indices)),
                "family_ids": "; ".join(sorted(str(data.loc[i, "family_id_norm"]) for i in indices)),
                "exact_match_fields": "; ".join(sorted(fields)),
            }
        )

    collapsed = data.loc[sorted(keep)].copy()
    return collapsed, pd.DataFrame(group_rows)


def style_workbook(path) -> None:
    wb = load_workbook(path)
    navy = "17365D"
    pale = "D9EAF7"
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            letter = col[0].column_letter
            width = min(42, max(11, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width
        if ws.title == "00_READ_ME":
            for cell in ws[2]:
                cell.fill = PatternFill("solid", fgColor=pale)
    wb.save(path)


def main() -> None:
    data = pd.read_excel(INPUT, sheet_name="04_Unblinded_Data")
    collapsed, groups = collapse_exact_duplicates(data)
    desc, comp = outcome_table(collapsed)
    readme = pd.DataFrame(
        [
            ["Purpose", "Exploratory sensitivity analysis after collapsing within-topic rows linked by identical First Claim or Abstract text."],
            ["Primary analysis", "The preregistered 87-row audit_results.xlsx remains primary."],
            ["Why added", "Duplicate source text was detected before opening the topic key; the sensitivity analysis was executed after unblinding."],
            ["Collapsed set", f"{len(collapsed)} rows: Topic 22 n={(collapsed.topic == 22).sum()}, Topic 25 n={(collapsed.topic == 25).sum()}."],
            ["Interpretation", "Robustness to exact text duplication only; not a replacement for an extended-family legal analysis."],
            ["Repeat assessment", "Round 2 was completed on the same day and is not a delayed reliability assessment."],
        ],
        columns=["Item", "Statement"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="00_READ_ME", index=False)
        desc.to_excel(writer, sheet_name="01_Descriptive", index=False)
        comp.to_excel(writer, sheet_name="02_Comparison", index=False)
        groups.to_excel(writer, sheet_name="03_Duplicate_Groups", index=False)
        collapsed[["record_id", "topic", "family_id_norm", *METRICS]].to_excel(
            writer, sheet_name="04_Collapsed_Data", index=False
        )
    style_workbook(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
