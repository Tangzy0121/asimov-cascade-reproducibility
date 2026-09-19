from __future__ import annotations

import hashlib
import json
import random
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(r"<project>/audit_workspace/topic22_25_single_reviewer_audit")
REPRESENTATIVES = Path(
    r"<project>/audit_workspace/family_id_audit/work/corrected/family_representatives.csv"
)
SOURCE = Path(r"<project>/Cascade\data\humanoid_safety_patents_v5_clean.xlsx")
SEED = 20260902
RETEST_N = 18

ENUMS = {
    "review_status": ["COMPLETE", "NEEDS_FULL_TEXT", "EXCLUDE_DATA_ERROR"],
    "safety_relevance": ["DIRECT", "PARTIAL", "INCIDENTAL", "NOT_SAFETY", "UNCLEAR"],
    "human_harm_link": ["DIRECT", "INDIRECT", "NONE", "UNCLEAR"],
    "contact_context": [
        "PHYSICAL_HRI", "COLLISION_CONTACT", "PROXIMITY_SEPARATION",
        "OBJECT_ENVIRONMENT", "NONE", "UNCLEAR",
    ],
    "sensing_type": ["FORCE_TORQUE", "PROXIMITY", "BOTH", "OTHER", "NONE", "UNCLEAR"],
    "sensor_evidence": ["YES", "NO", "UNCLEAR"],
    "decision_evidence": ["YES", "NO", "UNCLEAR"],
    "response_evidence": ["YES", "NO", "UNCLEAR"],
    "protective_action": [
        "STOP", "RETREAT", "SLOW", "LIMIT_FORCE", "WARNING",
        "ISOLATE_OR_FALLBACK", "MULTIPLE", "OTHER", "NONE", "UNCLEAR",
    ],
    "humanoid_scope": ["EXPLICIT_HUMANOID", "GENERAL_ROBOT", "NON_ROBOT", "UNCLEAR"],
    "barrier_evidence": ["YES", "NO", "UNCLEAR"],
    "propagation_evidence": ["EXPLICIT_CROSS_SUBSYSTEM", "PLAUSIBLE_ONLY", "ABSENT", "UNCLEAR"],
    "topic_scope": ["IN_SCOPE", "MIXED", "OFF_TOPIC", "UNCLEAR"],
    "evidence_location": ["FIRST_CLAIM", "ABSTRACT", "TITLE", "FULL_TEXT", "MULTIPLE", "UNCLEAR"],
    "reviewer_confidence": ["HIGH", "MEDIUM", "LOW"],
}

META_COLUMNS = [
    "record_id", "publication_no", "application_date", "applicant", "ipc",
    "title_english", "abstract_english", "first_claim_english",
]
CODING_COLUMNS = [
    "review_status", "safety_relevance", "human_harm_link", "contact_context",
    "sensing_type", "sensor_evidence", "decision_evidence", "response_evidence",
    "protective_action", "humanoid_scope", "barrier_evidence", "propagation_evidence",
    "topic_scope", "evidence_location", "exact_evidence_quote", "decision_reason",
    "reviewer_confidence", "review_minutes", "review_date",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def clean_text(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).replace("\x00", " ")
    return " ".join(text.split())


def load_records() -> pd.DataFrame:
    reps = pd.read_csv(REPRESENTATIVES, dtype={"Publication No": str, "family_id_norm": str})
    reps = reps.loc[reps["topic"].isin([22, 25])].copy()
    assert reps["topic"].value_counts().to_dict() == {22: 49, 25: 38}
    assert reps["family_id_norm"].nunique() == 87

    source = pd.read_excel(
        SOURCE,
        usecols=[
            "Publication No", "Application Date", "Applicant (Translation)", "IPC",
            "Title (English)", "Abstract (English)", "First Claim(English)",
        ],
        dtype={"Publication No": str},
    )
    source = source.drop_duplicates(subset=["Publication No"], keep="first")
    merged = reps.merge(source, on="Publication No", how="left", validate="one_to_one", indicator=True)
    assert (merged["_merge"] == "both").all(), "Some representative patents were not found in source data"
    merged = merged.drop(columns="_merge")

    rng = random.Random(SEED)
    order = list(range(len(merged)))
    rng.shuffle(order)
    merged = merged.iloc[order].reset_index(drop=True)
    merged["review_id"] = [f"R{i:03d}" for i in range(1, len(merged) + 1)]
    return merged


def reviewer_frame(records: pd.DataFrame, id_col: str) -> pd.DataFrame:
    out = pd.DataFrame()
    out["record_id"] = records[id_col]
    out["publication_no"] = records["Publication No"].map(clean_text)
    out["application_date"] = records["application_date"].map(clean_text)
    out["applicant"] = records["Applicant (Translation)"].map(clean_text)
    out["ipc"] = records["IPC"].map(clean_text)
    out["title_english"] = records["Title (English)"].map(clean_text)
    out["abstract_english"] = records["Abstract (English)"].map(clean_text)
    out["first_claim_english"] = records["First Claim(English)"].map(clean_text)
    for col in CODING_COLUMNS:
        out[col] = ""
    return out


def add_readme(ws, round_name: str, n_rows: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 104
    ws["A1"] = f"Topic 22/25 Structured Content Audit — {round_name}"
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws.merge_cells("A1:B1")
    rows = [
        ("Purpose", f"Review {n_rows} patent-family representatives while blinded to Topic identity and model outputs."),
        ("Edit", "Only pale-yellow cells in the REVIEW sheet. All coding cells start blank."),
        ("Do not open", "99_DO_NOT_OPEN_topic_key.xlsx until both rounds are complete."),
        ("Evidence order", "First Claim > Abstract > Title. Use FULL_TEXT only if you actually consult the original patent."),
        ("Uncertainty", "Choose UNCLEAR when evidence is insufficient. Do not guess."),
        ("Completion", "Set review_status=COMPLETE only after the evidence quote, reason, confidence, and all categorical fields are filled."),
        ("Example only", "A claim says external force is detected, compared with a threshold, then the joint stops: sensor=YES; decision=YES; response=YES; protective_action=STOP."),
        ("Terminology", "This is a single-reviewer structured content audit, not independent validation or double-blind review."),
    ]
    for row_idx, (label, value) in enumerate(rows, start=3):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
        ws.cell(row_idx, 1).font = Font(name="Arial", bold=True, color="17365D")
        ws.cell(row_idx, 2).font = Font(name="Arial", size=10)
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A13"] = "Legend"
    ws["A13"].font = Font(name="Arial", bold=True, color="17365D")
    ws["B13"] = "Blue/gray = source text (read only); pale yellow = reviewer input; red highlight = row marked COMPLETE but missing required fields."
    ws["B13"].alignment = Alignment(wrap_text=True)


def write_review_workbook(frame: pd.DataFrame, path: Path, round_name: str) -> None:
    wb = Workbook()
    readme = wb.active
    readme.title = "00_READ_ME"
    add_readme(readme, round_name, len(frame))

    ws = wb.create_sheet("01_REVIEW")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame)+1}"
    ws.row_dimensions[1].height = 42

    navy = PatternFill("solid", fgColor="17365D")
    source_fill = PatternFill("solid", fgColor="EAF1F8")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin, right=thin)

    for col_idx, col in enumerate(frame.columns, start=1):
        cell = ws.cell(1, col_idx, col)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = navy
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_idx, row in enumerate(frame.itertuples(index=False), start=2):
        ws.row_dimensions[row_idx].height = 72
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if col_idx <= len(META_COLUMNS):
                cell.fill = source_fill
                cell.protection = Protection(locked=True)
            else:
                cell.fill = input_fill
                cell.protection = Protection(locked=False)

    widths = {
        "record_id": 11, "publication_no": 17, "application_date": 14, "applicant": 25,
        "ipc": 18, "title_english": 34, "abstract_english": 58, "first_claim_english": 68,
        "review_status": 20, "safety_relevance": 18, "human_harm_link": 18,
        "contact_context": 23, "sensing_type": 18, "sensor_evidence": 17,
        "decision_evidence": 18, "response_evidence": 18, "protective_action": 22,
        "humanoid_scope": 21, "barrier_evidence": 18, "propagation_evidence": 27,
        "topic_scope": 15, "evidence_location": 18, "exact_evidence_quote": 48,
        "decision_reason": 44, "reviewer_confidence": 19, "review_minutes": 15,
        "review_date": 15,
    }
    for idx, col in enumerate(frame.columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 15)

    for col, values in ENUMS.items():
        col_idx = frame.columns.get_loc(col) + 1
        col_letter = get_column_letter(col_idx)
        formula = '"' + ",".join(values) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Select a value from the dropdown list."
        dv.errorTitle = "Invalid code"
        dv.prompt = "Use CODEBOOK.md; choose UNCLEAR rather than guessing."
        dv.promptTitle = col
        dv.showErrorMessage = True
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{len(frame)+1}")

    minutes_col = get_column_letter(frame.columns.get_loc("review_minutes") + 1)
    dv_minutes = DataValidation(type="decimal", operator="between", formula1="0", formula2="240", allow_blank=True)
    ws.add_data_validation(dv_minutes)
    dv_minutes.add(f"{minutes_col}2:{minutes_col}{len(frame)+1}")

    status_col = get_column_letter(frame.columns.get_loc("review_status") + 1)
    last_col = get_column_letter(len(frame.columns))
    required_start = get_column_letter(frame.columns.get_loc("safety_relevance") + 1)
    required_end = get_column_letter(frame.columns.get_loc("reviewer_confidence") + 1)
    red_fill = PatternFill("solid", fgColor="F4CCCC")
    ws.conditional_formatting.add(
        f"A2:{last_col}{len(frame)+1}",
        FormulaRule(
            formula=[f'AND(${status_col}2="COMPLETE",COUNTBLANK(${required_start}2:${required_end}2)>0)'],
            fill=red_fill,
        ),
    )

    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.password = "audit"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)


def write_key(records: pd.DataFrame, retest: pd.DataFrame, path: Path) -> None:
    key = records[[
        "review_id", "family_id_norm", "topic", "Publication No", "source_row", "family_doc_id"
    ]].copy()
    key = key.merge(
        retest[["review_id", "retest_id"]], on="review_id", how="left", validate="one_to_one"
    )
    key = key.rename(columns={"Publication No": "publication_no"})
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        key.to_excel(writer, sheet_name="DO_NOT_OPEN", index=False)
        meta = pd.DataFrame(
            {
                "field": ["warning", "random_seed", "round1_n", "retest_n", "topic_22_n", "topic_25_n"],
                "value": [
                    "Do not open until both review rounds are complete.", SEED, len(records),
                    len(retest), int((records.topic == 22).sum()), int((records.topic == 25).sum()),
                ],
            }
        )
        meta.to_excel(writer, sheet_name="MANIFEST", index=False)
    wb = __import__("openpyxl").load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C00000")
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 24
    wb.save(path)


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    records = load_records()
    round1 = reviewer_frame(records, "review_id")

    rng = random.Random(SEED + 1)
    selected_ids = rng.sample(records["review_id"].tolist(), RETEST_N)
    retest = records.loc[records["review_id"].isin(selected_ids)].copy()
    retest_order = list(range(len(retest)))
    rng.shuffle(retest_order)
    retest = retest.iloc[retest_order].reset_index(drop=True)
    retest["retest_id"] = [f"T{i:03d}" for i in range(1, RETEST_N + 1)]
    round2 = reviewer_frame(retest, "retest_id")

    round1_path = ROOT / "01_round1_blinded_review.xlsx"
    round2_path = ROOT / "02_round2_delayed_retest.xlsx"
    key_path = ROOT / "99_DO_NOT_OPEN_topic_key.xlsx"
    write_review_workbook(round1, round1_path, "Round 1")
    write_review_workbook(round2, round2_path, "Delayed retest")
    write_key(records, retest, key_path)

    manifest = {
        "seed": SEED,
        "retest_n": RETEST_N,
        "source_files": {str(REPRESENTATIVES): sha256(REPRESENTATIVES), str(SOURCE): sha256(SOURCE)},
        "outputs": {str(p.name): sha256(p) for p in [round1_path, round2_path, key_path]},
        "counts": {"round1": len(round1), "round2": len(round2), "topic_22": 49, "topic_25": 38},
    }
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
