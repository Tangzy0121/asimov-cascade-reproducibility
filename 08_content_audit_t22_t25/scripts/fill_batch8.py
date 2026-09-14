"""Write batch 8 (rows 86-88, R085-R087) reviewer fields + final full validation."""

import datetime
import sys
from collections import Counter

import openpyxl

PATH = "01_round1_blinded_review.xlsx"
TODAY = datetime.date.today().isoformat()

ROWS = {
   86: [
      "COMPLETE",
      "DIRECT",
      "INDIRECT",
      "COLLISION_CONTACT",
      "FORCE_TORQUE",
      "YES",
      "YES",
      "YES",
      "STOP",
      "GENERAL_ROBOT",
      "YES",
      "ABSENT",
      "IN_SCOPE",
      "FIRST_CLAIM",
      "the stop instruction by the said detecting part detects the physical quantity with a first threshold value and…second threshold value…respectively comparing…stops, and the physical quantity is more than the second threshold value, the robot than stopping method of the predetermined short time interval",
      "人协调机器人共享作业空间：检测随接触力变化的物理量→与第一/第二阈值分别比较→超第一阈值按预定方式停止、超第二阈值以更短時間间隔急停（分级 STOP）。完整机制链，接触力涉人共享空间但未明言伤害，harm=INDIRECT。",
      "HIGH",
      4,
      TODAY,
   ],
   87: [
      "COMPLETE",
      "DIRECT",
      "DIRECT",
      "PHYSICAL_HRI",
      "FORCE_TORQUE",
      "YES",
      "YES",
      "YES",
      "MULTIPLE",
      "GENERAL_ROBOT",
      "YES",
      "ABSENT",
      "IN_SCOPE",
      "FIRST_CLAIM",
      "when a detected value of said force sensor exceeds a predetermined value, stopping said robot or controlling operation of said robot so that a detected value of said force sensor becomes smaller…a limiter which limits a work area of said human so as to prevent contact by said human with said first robot portion",
      "人机共享区域：力传感器检测值超限→停止或控制使力减小（STOP+LIMIT_FORCE），另设限制器限制人的工作区域以防止人与机器人第一部份接触（ISOLATE），action=MULTIPLE。明确防止人体接触，harm=DIRECT。",
      "HIGH",
      5,
      TODAY,
   ],
   88: [
      "COMPLETE",
      "DIRECT",
      "INDIRECT",
      "OBJECT_ENVIRONMENT",
      "FORCE_TORQUE",
      "YES",
      "YES",
      "YES",
      "ISOLATE_OR_FALLBACK",
      "NON_ROBOT",
      "YES",
      "ABSENT",
      "IN_SCOPE",
      "FIRST_CLAIM",
      "如果所监测的操作参数偏离所需状态，则通过自给的第二辅助真空源(20, 26)在抽吸夹持器(12)中产生辅助负压",
      "人机协作真空搬运装置：搬运中监测状态变量（真空压力类，sensing=FORCE_TORQUE）→偏离所需状态时自给辅助真空源产生辅助负压保持工件（故障应急保持，ISOLATE_OR_FALLBACK），防掉件的故障响应是主要安全功能（DIRECT）。掉件伤人路径合理但未明言，harm=INDIRECT。主体为真空搬运装置而非机器人系统，NON_ROBOT。",
      "MEDIUM",
      5,
      TODAY,
   ],
}

LEGAL = {
   9: {"COMPLETE", "NEEDS_FULL_TEXT", "EXCLUDE_DATA_ERROR"},
   10: {"DIRECT", "PARTIAL", "INCIDENTAL", "NOT_SAFETY", "UNCLEAR"},
   11: {"DIRECT", "INDIRECT", "NONE", "UNCLEAR"},
   12: {
      "PHYSICAL_HRI",
      "COLLISION_CONTACT",
      "PROXIMITY_SEPARATION",
      "OBJECT_ENVIRONMENT",
      "NONE",
      "UNCLEAR",
   },
   13: {"FORCE_TORQUE", "PROXIMITY", "BOTH", "OTHER", "NONE", "UNCLEAR"},
   14: {"YES", "NO", "UNCLEAR"},
   15: {"YES", "NO", "UNCLEAR"},
   16: {"YES", "NO", "UNCLEAR"},
   17: {
      "STOP",
      "RETREAT",
      "SLOW",
      "LIMIT_FORCE",
      "WARNING",
      "ISOLATE_OR_FALLBACK",
      "MULTIPLE",
      "OTHER",
      "NONE",
      "UNCLEAR",
   },
   18: {"EXPLICIT_HUMANOID", "GENERAL_ROBOT", "NON_ROBOT", "UNCLEAR"},
   19: {"YES", "NO", "UNCLEAR"},
   20: {"EXPLICIT_CROSS_SUBSYSTEM", "PLAUSIBLE_ONLY", "ABSENT", "UNCLEAR"},
   21: {"IN_SCOPE", "MIXED", "OFF_TOPIC", "UNCLEAR"},
   22: {"FIRST_CLAIM", "ABSTRACT", "TITLE", "FULL_TEXT", "MULTIPLE", "UNCLEAR"},
   25: {"HIGH", "MEDIUM", "LOW"},
}

wb = openpyxl.load_workbook(PATH)
ws = wb["01_REVIEW"]
for r, vals in ROWS.items():
   for i, v in enumerate(vals):
      ws.cell(row=r, column=9 + i, value=v)
try:
   wb.save(PATH)
except PermissionError:
   print("LOCKED: file still open in Excel, nothing saved")
   sys.exit(2)
print("SAVED rows 86-88")

# ---- final full validation ----
wb = openpyxl.load_workbook(PATH)
ws = wb["01_REVIEW"]
bk = openpyxl.load_workbook("01_round1_blinded_review.backup.xlsx")["01_REVIEW"]
diff = sum(
   1
   for r in range(1, 89)
   for c in range(1, 9)
   if ws.cell(r, c).value != bk.cell(r, c).value
)
print("A-H diffs vs backup:", diff)
problems = []
for r in range(2, 89):
   for c in range(9, 28):
      v = ws.cell(r, c).value
      if v in (None, ""):
         problems.append((r, c, "EMPTY"))
      elif c in LEGAL and v not in LEGAL[c]:
         problems.append((r, c, "ILLEGAL:" + str(v)))
print("field problems:", problems if problems else "NONE")
status = Counter(ws.cell(r, 9).value for r in range(2, 89))
print("status:", dict(status))
print("safety:", dict(Counter(ws.cell(r, 10).value for r in range(2, 89))))
print("scope:", dict(Counter(ws.cell(r, 21).value for r in range(2, 89))))
print("conf:", dict(Counter(ws.cell(r, 25).value for r in range(2, 89))))
print("propagation:", dict(Counter(ws.cell(r, 20).value for r in range(2, 89))))
unclear_cells = sum(
   1 for r in range(2, 89) for c in range(9, 28) if ws.cell(r, c).value == "UNCLEAR"
)
unclear_rows = sorted(
   str(ws.cell(r, 1).value)
   for r in range(2, 89)
   if any(ws.cell(r, c).value == "UNCLEAR" for c in range(9, 28))
)
print("rows containing UNCLEAR:", unclear_rows)
print("total UNCLEAR cells:", unclear_cells)
chain = sum(
   1 for r in range(2, 89) if all(ws.cell(r, c).value == "YES" for c in (14, 15, 16))
)
print("full mechanism chains (sensor+decision+response all YES):", chain)
