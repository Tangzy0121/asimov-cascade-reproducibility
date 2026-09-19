#!/usr/bin/env python
# -*- coding: utf-8 -*-
# File: scripts/build_dual_review_package.py
# Purpose : Build the 77-topic dual human review workbook for the ICRA 2027
#           revision (reviewer asked for full human adjudication of all
#           non-noise topics + human Cohen's kappa, replacing the
#           "57 unadjudicated topics" limitation). One row per topic with
#           label / top keywords / up to 3 representative patent excerpts,
#           followed by two independent EMPTY rater blocks (Rater A / B).
# Inputs  : output/time_analysis/04_bertopic_time/topic_summary.csv
#           output/time_analysis/10_cascade/llm_safety_labels.csv  (set check only)
#           ../data/humanoid_safety_patents_v5_clean.xlsx          (patent numbers)
# Outputs : output/dual_review/dual_review_workbook.xlsx
#           output/dual_review/label_definitions.md
# Integrity: NO LLM label and NO human judgment is written into the rater
#            columns — the LLM labels file is only used to assert that the
#            topic set matches the frozen screening set (77 non-noise topics).
#            Label schema mirrors the P1 pilot file
#            output/subsystem_validation/human_review/P1_topic_safety_reviewed.csv
# Usage   : <python>/envs/<env>/python.exe -X utf8 scripts/build_dual_review_package.py
from __future__ import annotations

import ast
import csv
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TOPIC_SUMMARY = (
    PROJECT_ROOT / "output" / "time_analysis" / "04_bertopic_time" / "topic_summary.csv"
)
LLM_LABELS = (
    PROJECT_ROOT / "output" / "time_analysis" / "10_cascade" / "llm_safety_labels.csv"
)
PATENTS_XLSX = (
    PROJECT_ROOT.parent / "data" / "humanoid_safety_patents_v5_clean.xlsx"
)
OUT_DIR = PROJECT_ROOT / "output" / "dual_review"
OUT_XLSX = OUT_DIR / "dual_review_workbook.xlsx"
OUT_LABEL_DEFS = OUT_DIR / "label_definitions.md"

EXCERPT_CHARS = 400          # representative abstract truncation length
MAX_REP_DOCS = 3
TOP_N_KEYWORDS = 10

# Label schema — value sets confirmed against
# output/subsystem_validation/human_review/P1_topic_safety_reviewed.csv
SAFETY_JUDGMENT = ["DIRECT", "PARTIAL", "INCIDENTAL", "NOT_SAFETY", "UNCLEAR"]
HARM_LINK = ["DIRECT", "INDIRECT", "UNCLEAR"]
CASCADE_ROLE = ["SAFETY_BARRIER", "PROPAGATION_NODE", "CONTEXT_ONLY",
                "OUT_OF_SCOPE", "UNCLEAR"]
CONFIDENCE = ["HIGH", "MEDIUM", "LOW"]

# Workbook columns: topic info + representative evidence + rater blocks
INFO_COLS = ["topic_id", "topic_label", "doc_count", "top_keywords"]
REP_COLS = [f"rep{i}_{part}" for i in (1, 2, 3) for part in ("patent", "excerpt")]
RATER_FIELDS = ["safety_judgment", "harm_link", "cascade_role",
                "scope_limitation", "confidence", "note"]
HEADERS = (INFO_COLS + REP_COLS
           + [f"A_{f}" for f in RATER_FIELDS]
           + [f"B_{f}" for f in RATER_FIELDS])

# enum field -> allowed values (None = free text, no dropdown)
ENUMS = {
    "safety_judgment": SAFETY_JUDGMENT,
    "harm_link": HARM_LINK,
    "cascade_role": CASCADE_ROLE,
    "scope_limitation": None,
    "confidence": CONFIDENCE,
    "note": None,
}


def load_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def norm_text(s: object) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def build_abstract_index() -> dict[str, str]:
    """normalized clean_abstract -> Publication No."""
    df = pd.read_excel(PATENTS_XLSX, usecols=["Publication No", "clean_abstract"])
    idx = {}
    for _, r in df.iterrows():
        key = norm_text(r["clean_abstract"])
        if key and key != "nan":
            idx[key] = str(r["Publication No"]).strip()
    return idx


def truncate(text: str, limit: int = EXCERPT_CHARS) -> str:
    text = re.sub(r"\s+", " ", str(text)).strip()
    return text if len(text) <= limit else text[:limit].rstrip() + " …"


def collect_topics() -> list[dict]:
    """One record per non-noise topic, with up to 3 representative patents."""
    summary = load_csv(TOPIC_SUMMARY)
    llm = load_csv(LLM_LABELS)
    topics = [r for r in summary if r["Topic"] != "-1"]
    llm_ids = {int(r["bertopic_id"]) for r in llm}
    topic_ids = {int(r["Topic"]) for r in topics}
    assert topic_ids == llm_ids, (
        f"topic set mismatch: summary {len(topic_ids)} vs llm_labels {len(llm_ids)}; "
        f"diff={topic_ids ^ llm_ids}"
    )
    abs2pn = build_abstract_index()

    records = []
    for r in sorted(topics, key=lambda x: int(x["Topic"])):
        docs = ast.literal_eval(r["Representative_Docs"])[:MAX_REP_DOCS]
        reps = []
        for d in docs:
            key = norm_text(d)
            pn = abs2pn.get(key, "")
            if not pn:  # prefix fallback (truncation differences)
                pn = next((v for k, v in abs2pn.items()
                           if k[:150] == key[:150]), "")
            reps.append((pn, truncate(d)))
        while len(reps) < MAX_REP_DOCS:
            reps.append(("", ""))
        words = ast.literal_eval(r["Representation"])[:TOP_N_KEYWORDS]
        records.append({
            "topic_id": int(r["Topic"]),
            "topic_label": r["Name"],
            "doc_count": int(r["Count"]),
            "top_keywords": ", ".join(words),
            "reps": reps,
        })
    return records


def write_workbook(records: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "dual_review"

    header_fill = PatternFill("solid", start_color="FFD9E1F2")   # info: blue
    rep_fill = PatternFill("solid", start_color="FFE2EFDA")      # evidence: green
    a_fill = PatternFill("solid", start_color="FFFFF2CC")        # rater A: yellow
    b_fill = PatternFill("solid", start_color="FFFCE4D6")        # rater B: orange
    n_info, n_rep = len(INFO_COLS), len(REP_COLS)

    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True)
        if c <= n_info:
            cell.fill = header_fill
        elif c <= n_info + n_rep:
            cell.fill = rep_fill
        elif c <= n_info + n_rep + len(RATER_FIELDS):
            cell.fill = a_fill
        else:
            cell.fill = b_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for i, rec in enumerate(records, start=2):
        row = ([rec["topic_id"], rec["topic_label"], rec["doc_count"],
                rec["top_keywords"]]
               + [v for rep in rec["reps"] for v in rep]
               + [""] * (2 * len(RATER_FIELDS)))
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # dropdown validations for enum rater fields, applied to all data rows
    last_row = len(records) + 1
    for prefix in ("A", "B"):
        for field in RATER_FIELDS:
            allowed = ENUMS[field]
            if not allowed:
                continue
            col = HEADERS.index(f"{prefix}_{field}") + 1
            letter = get_column_letter(col)
            dv = DataValidation(
                type="list", formula1='"' + ",".join(allowed) + '"',
                allow_blank=True, showErrorMessage=True,
                errorTitle="invalid label",
                error=f"choose one of: {', '.join(allowed)}",
            )
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{last_row}")

    # column widths
    widths = {"topic_id": 9, "topic_label": 34, "doc_count": 9,
              "top_keywords": 46}
    for i in (1, 2, 3):
        widths[f"rep{i}_patent"] = 16
        widths[f"rep{i}_excerpt"] = 60
    for prefix in ("A", "B"):
        widths[f"{prefix}_safety_judgment"] = 16
        widths[f"{prefix}_harm_link"] = 12
        widths[f"{prefix}_cascade_role"] = 20
        widths[f"{prefix}_scope_limitation"] = 34
        widths[f"{prefix}_confidence"] = 12
        widths[f"{prefix}_note"] = 30
    for c, name in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)

    ws.freeze_panes = "E2"   # first row + id/label/count/keywords stay visible
    ws.row_dimensions[1].height = 30
    wb.save(OUT_XLSX)


LABEL_DEFS_MD = """# Dual-Rater Label Definitions (label schema)

This file defines the meaning of the values allowed in the Rater A / Rater B
coding columns of `dual_review_workbook.xlsx`. The value sets are identical to
the P1 pilot coding
(`output/subsystem_validation/human_review/P1_topic_safety_reviewed.csv`) and
correspond to the safety-cascade interpretation framework in Section IV-C of
the paper. Read this file in full before coding; if you are unsure about a
topic, record the issue in the `*_note` column per the protocol instead of
leaving the cell blank.

## 1. safety_judgment (topic safety relevance)

| Value | Definition |
| ----- | ---------- |
| `DIRECT` | The patents in the topic **explicitly** aim to protect people or prevent human injury, or directly describe the safety risks a humanoid robot poses to people together with the corresponding mechanisms. |
| `PARTIAL` | The topic contains explicit safety mechanisms (e.g., balance recovery, safety braking, collision monitoring), but safety is not the main contribution of most patents in the topic, or the topic clearly mixes safety and non-safety content. |
| `INCIDENTAL` | Safety content appears only as an ancillary condition or indirectly (e.g., obstacle avoidance performed to complete a surveying or manipulation task); protecting people is not the goal of the invention. |
| `NOT_SAFETY` | The evidence text contains no substantive safety mechanism; "safety" keywords are background or stylistic wording (e.g., "human-like" describing a driving style). |
| `UNCLEAR` | The evidence is insufficient, or the topic is severely mixed (e.g., grippers, electrodes, and measurement lumped into one topic), so no reliable judgment can be made. |

## 2. harm_link (link to human injury)

| Value | Definition |
| ----- | ---------- |
| `DIRECT` | The patent text directly describes human injury, human-contact risk, or people as the object of protection (e.g., pinched fingers, contact in shared spaces). |
| `INDIRECT` | Human injury is **inferred across scenarios** from the evidence (e.g., a fall of a quadruped or wheel-legged platform could affect nearby people); the text does not directly report injuries to people. |
| `UNCLEAR` | The harm link cannot be determined. |

## 3. cascade_role (role of the topic in the Asimov Cascade)

| Value | Definition |
| ----- | ---------- |
| `SAFETY_BARRIER` | The topic is a safety barrier that blocks or mitigates cascading failures (emergency stop, safety braking, shutdown after collision detection, balance recovery, etc.). |
| `PROPAGATION_NODE` | The topic describes the information/control transmission chain across perception, planning, control, and actuation; interface inconsistencies may propagate along the chain into consequences such as instability or collision. |
| `CONTEXT_ONLY` | Safety is merely a background condition for accomplishing other tasks; the topic can serve only as a context chain, not as core safety evidence. |
| `OUT_OF_SCOPE` | The topic does not primarily belong to humanoid robotics research (e.g., autonomous-vehicle lane changing, tobacco warehouse equipment) and should be excluded from the main analysis. |
| `UNCLEAR` | The role cannot be determined. |

## 4. scope_limitation (scope limitation, free text)

Record the applicability limits of the topic's evidence; there is no dropdown
and the entry is free-form. Common phrasings (following actual P1 usage):

- Non-humanoid evidence platform: "All evidence comes from wheel-legged or quadruped robots, not humanoids; human injury is only an indirect inference."
- Mixed topic: "The topic mixes X, Y, and a small number of humanoid-robot patents."
- Nature of the evidence: "The patents describe risk-prevention methods, not records of real accidents."

If there is no scope issue, enter `none` or a brief note — **do not
leave the cell blank**.

## 5. confidence (coding confidence)

| Value | Definition |
| ----- | ---------- |
| `HIGH` | The evidence is direct and the basis for the judgment is explicit (e.g., the text explicitly describes human contact and a shutdown mechanism). |
| `MEDIUM` | The basis for the judgment holds but requires cross-scenario inference, or the topic is somewhat mixed. |
| `LOW` | The evidence is thin or the topic is severely mixed; the judgment is for adjudication reference only. |

## 6. note (coding note, free text)

Record the rationale, concerns, and suggestions (e.g., "recommend splitting
this topic"). This complements scope_limitation: scope_limitation states "how
far the evidence can be extrapolated", while note states "why I judged it
this way".
"""


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    records = collect_topics()
    write_workbook(records)
    OUT_LABEL_DEFS.write_text(LABEL_DEFS_MD, encoding="utf-8")
    print(f"topics written : {len(records)}")
    print(f"workbook       : {OUT_XLSX}")
    print(f"label defs     : {OUT_LABEL_DEFS}")


if __name__ == "__main__":
    main()
